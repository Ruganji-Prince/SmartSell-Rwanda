"""Reproducible SAS annual-table extraction, with source cell provenance.

These are published sector aggregates, never individual Tomato/Carrot records.
Future microdata adapters should emit separately labelled crop-specific tables.
"""
import math
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from .utilities import ROOT, save_json, sha256

DISTRICTS = "Nyarugenge Gasabo Kicukiro Nyanza Gisagara Nyaruguru Huye Nyamagabe Ruhango Muhanga Kamonyi Karongi Rutsiro Rubavu Nyabihu Ngororero Rusizi Nyamasheke Rulindo Gakenke Musanze Burera Gicumbi Rwamagana Nyagatare Gatsibo Kayonza Kirehe Ngoma Bugesera".split()


def number(value):
    if isinstance(value, (int, float)) and math.isfinite(value):
        return float(value)
    # Dashes and empty cells remain missing, not invented zeroes.
    return None


def extract_annual_tables(path, output=None):
    output = output or ROOT / "data/processed/nisr"
    output.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(path, data_only=True, read_only=True)
    district_rows, usage_rows, sowing_rows, quarantine = [], [], [], []
    table_map = {13: ("A", "harvested_area_ha"), 14: ("B", "harvested_area_ha"),
                 15: ("C", "harvested_area_ha"), 16: ("A", "yield_kg_ha"),
                 17: ("B", "yield_kg_ha"), 18: ("C", "yield_kg_ha"),
                 21: ("A", "production_mt"), 22: ("B", "production_mt"), 23: ("C", "production_mt")}
    for n, (season, indicator) in table_map.items():
        sheet = workbook[f"Table {n}"]
        rows = list(sheet.values)
        header = next(i for i, r in enumerate(rows) if any(str(v).strip() in ["District", "District/Crop"] for v in r))
        district_col = next(j for j, v in enumerate(rows[header]) if str(v).strip() in ["District", "District/Crop"])
        for row_num, row in enumerate(rows[header + 1:], header + 2):
            district = str(row[district_col]).strip().title()
            if district not in DISTRICTS:
                continue
            for col in range(district_col + 1, len(row)):
                crop = rows[header][col]
                if crop is None:
                    continue
                crop = str(crop).strip()
                value = number(row[col])
                record = dict(year=2024, season=season, district=district, crop_group=crop,
                              indicator=indicator, value=value, table=n,
                              source_cell=f"{get_column_letter(col+1)}{row_num}")
                if value is not None and value < 0:
                    quarantine.append(record)
                else:
                    district_rows.append(record)
    for n, season in [(24, "A"), (25, "B"), (26, "C")]:
        rows = list(workbook[f"Table {n}"].values)
        hi = next(i for i,r in enumerate(rows) if "Crops" in r)
        ci = rows[hi].index("Crops")
        for rn, row in enumerate(rows[hi + 1:], hi + 2):
            crop = row[ci]
            if not crop or str(crop).startswith("Source"):
                continue
            for col in range(ci+1, len(row)):
                if rows[hi][col] is None:
                    continue
                value = number(row[col])
                rec = dict(year=2024, season=season, geography="Rwanda (national)",
                           crop_group=str(crop).strip(), use=str(rows[hi][col]).strip(), percent=value,
                           denominator="Published share of crop-group production; not stored quantity or daily exposure",
                           table=n, source_cell=f"{get_column_letter(col+1)}{rn}")
                if value is not None and not 0 <= value <= 100:
                    quarantine.append(rec)
                else:
                    usage_rows.append(rec)
    for n in [27, 28]:
        rows = list(workbook[f"Table {n}"].values)
        hi = next(i for i,r in enumerate(rows) if "District" in r or "Crops" in r)
        ci = 1
        for rn, row in enumerate(rows[hi+1:], hi+2):
            entity = row[ci]
            if not entity or str(entity).startswith("Source"):
                continue
            if n == 27 and str(entity).strip().title() not in DISTRICTS:
                continue
            for col in range(2, min(len(row), 23)):
                value = number(row[col])
                sowing_rows.append(dict(year=2024, entity=str(entity).strip(),
                                       entity_type="district" if n == 27 else "crop_group",
                                       season="A" if col < 9 else "B" if col < 18 else "C",
                                       sowing_window=str(rows[hi][col]).strip(), percent=value,
                                       table=n, source_cell=f"{get_column_letter(col+1)}{rn}"))
    district = pd.DataFrame(district_rows)
    usage = pd.DataFrame(usage_rows)
    district.to_csv(output / "district_indicators.csv", index=False)
    usage.to_csv(output / "production_use.csv", index=False)
    pd.DataFrame(sowing_rows).to_csv(output / "sowing_windows.csv", index=False)
    veg = district.loc[district.crop_group.str.lower().eq("vegetables")].copy()
    veg["crop_group"] = "Vegetables"
    veg.to_csv(output / "vegetable_district_context.csv", index=False)
    pd.DataFrame(quarantine, columns=["year", "season", "district", "crop_group", "indicator",
                 "value", "percent", "table", "source_cell"]).to_csv(output / "quarantine.csv", index=False)
    save_json(output / "extraction_metadata.json", dict(workbook=path.name, sha256=sha256(path),
              tables=sorted(list(table_map)+[24,25,26,27,28]), district_rows=len(district),
              vegetable_rows=len(veg), usage_rows=len(usage), quarantine_rows=len(quarantine),
              missing_values="Blank/dash cells remain missing; no daily deterioration inferred"))
    workbook.close()
    return veg, usage


def attach_context(prices, vegetables):
    """Attach historical all-season district context, without inventing harvest season.

    Seasons overlap and a market date does not identify the production season.
    Use 2024 district totals (sum production/area; area-weighted yield derived).
    Official seasonal yields remain in the source table, separate from this ratio.
    """
    pivot = vegetables.pivot(index=["district", "season"], columns="indicator", values="value").reset_index()
    context = pivot.groupby("district")[["production_mt", "harvested_area_ha"]].sum(min_count=1).reset_index()
    context["yield_kg_ha"] = 1000 * context.production_mt / context.harvested_area_ha.where(context.harvested_area_ha.gt(0))
    context = context.rename(columns={c:f"nisr_2024_vegetables_{c}" for c in context if c != "district"})
    joined = prices.merge(context, left_on="region", right_on="district", how="left", validate="many_to_one")
    joined["nisr_context_year"] = 2024
    mapping = joined[["region", "district"]].drop_duplicates().sort_values("region")
    mapping.to_csv(ROOT / "data/processed/nisr/district_mapping.csv", index=False)
    return joined
